"""Read the supplied editable grammar as data, preserving tabs and source blocks.
Usage: bundled-python extract-kasem-grammar.py <source.docx>
No sentence is generated from dictionary words. Every pair has a source block.
"""
import sys, json, re, hashlib, zipfile
from xml.etree import ElementTree as ET
from PIL import Image
from io import BytesIO
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
STARTS = [(123,1,'Sounds and writing'),(196,2,'Greetings and useful phrases'),(290,3,'Clause structure'),(437,4,'Noun phrases and pronouns'),(639,5,'Place time and adverb phrases'),(704,6,'Joining clauses'),(880,7,'Noun classes'),(1033,8,'Verb phrases'),(1476,9,'Small words and questions')]

def text(element):
    return ''.join(n.text or '' if n.tag == W+'t' else '\t' if n.tag == W+'tab' and not n.attrib else '\n' if n.tag == W+'br' else '' for n in element.iter())

def clean(value):
    return re.sub(r'\s+',' ',value).strip(' \t.,;:‘’“”\"-')

def chapter(index):
    return next(((number,title) for start,number,title in reversed(STARTS) if index >= start),(0,'Front matter'))

def valid(form):
    return bool(form and re.fullmatch(r'[A-Za-zɛɔŋƐƆŊáéíóúàèìòùÁÉÍÓÚÀÈÌÒÙ\s?!.,’\-]+',form)) and not re.search(r'\b(?:the|this|class|before|negative|particle|compare|habitual|in progress|or|tone|also|and)\b',form,re.I)

def main(path):
    source=Path(path); doc=Document(source); blocks=[]; figures=[]
    out=ROOT/'data/grammar-book-seed';out.mkdir(parents=True,exist_ok=True)
    image_out=out/'images';image_out.mkdir(exist_ok=True)
    descriptions=['People at a market','A walking woman','Three seated men','A child in profile',
        'A thatched house','The sun above the horizon','A chain with connected links','Two bowls',
        'A seated woman holding a utensil','A woman cooking over a fire',
        'A child, mouth, room and bowl','Two dogs','Two woven fans','Two hoes',
        'A cropped figure of a woman','Two figures in motion','People cooking in a courtyard',
        'A woman tending a cooking pot','A seated man working with his hands']
    for index,element in enumerate(doc.element.body):
        for node in element.iter():
            if not node.tag.endswith('}blip'):continue
            rel=node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            if not rel:continue
            part=doc.part.rels[rel].target_part;name=Path(str(part.partname)).name
            raw=part.blob;(image_out/name).write_bytes(raw);width,height=Image.open(BytesIO(raw)).size
            number=int(re.search(r'\d+',name)[0])
            figures.append({'file':name,'block':index,'chapter':chapter(index)[0],
                'description':descriptions[number-1],'width':width,'height':height,
                'sourceSha256':hashlib.sha256(raw).hexdigest(),'credit':'Original source illustration, A Basic Grammar of Kasem, GILLBT.'})
    for index,element in enumerate(doc.element.body):
        if element.tag not in [W+'p',W+'tbl']: continue
        number,title=chapter(index)
        blocks.append({'index':index,'chapter':number,'section':title,'kind':'table' if element.tag==W+'tbl' else 'paragraph','text':text(element),'rows':[[text(c) for c in row.findall(W+'tc')] for row in element.findall(W+'tr')] if element.tag==W+'tbl' else [],'figures':[f['file'] for f in figures if f['block']==index]})
    by_index={b['index']:b for b in blocks}; entries={}; examples={}; withheld=[]
    ref=lambda index:f'Chapter {chapter(index)[0]}, DOCX block {index}'
    def pair(form,meaning,index,pos='Not specified in source',note='',example=False,gloss=''):
        original=form; form=clean(form); meaning=clean(meaning)
        if re.search(r'[ɛɔŋ]|\b(?:tg|sice|wice|peat|didnotsee)\b',meaning,re.I):
            withheld.append({'sourceRef':ref(index),'form':original,'meaning':meaning,'reason':'Damaged or mixed-language gloss; retained in chapter reference.'});return
        if not valid(form) or not meaning or '�' in meaning:
            withheld.append({'sourceRef':ref(index),'form':original,'meaning':meaning,'reason':'Damaged or incomplete transcription; retained in chapter reference.'});return
        key=(form.lower(),meaning.lower(),pos if not example else '')
        target=examples if example else entries
        if key in target:
            if ref(index) not in target[key]['sourceRefs']:target[key]['sourceRefs'].append(ref(index))
            return
        if example:
            target[key]={'kasem':form,'english':meaning,'sourceRefs':[ref(index)],'sourceRef':ref(index),'section':chapter(index)[1],'note':note,'sourceGloss':gloss,'translationKind':'source-gloss' if gloss else 'printed-translation'}
        else:
            target[key]={'id':'gillbt83-'+hashlib.sha256('|'.join(key).encode()).hexdigest()[:24],'headword':form,'translation':meaning,'partOfSpeech':pos,'alternateForms':[],'sourceRefs':[ref(index)],'sourceForm':original,'note':note}
    # Explicit example/lexical lines, excluding explanatory English and headings.
    inline=set([141,142,146,147,150,151,152,158,159,160,161,165,166,167,168,181,182,183,184,185,186,187,188,189,190,270,271,272,273,274,275,276,279,280,281,282,283,284,285,286,287,288,303,332,335,336,338,339,390,392,394,418,442,443,503,505,512,513,520,522,524,526,528,532,550,555,556,558,560,562,576,582,586,588,596,599,610,615,616,617,621,623,624,625,631,636,648,651,654,656,661,663,665,668,669,672,674,677,679,682,690,699,765,790,798,804,815,816,857,862,873,910,1054,1056,1074,1076,1078,1082,1084,1097,1098,1099,1100,1101,1102,1105,1106,1108,1111,1112,1113,1116,1117,1121,1122,1123,1125,1126,1127,1128,1130,1131,1133,1134,1135,1137,1138,1143,1144,1145,1146,1147,1149,1150,1161,1162,1163,1171,1172,1173,1174,1176,1177,1179,1180,1181,1205,1206,1208,1209,1220,1222,1228,1229,1230,1231,1234,1239,1241,1244,1246,1248,1251,1252,1253,1256,1260,1261,1263,1265,1266,1268,1269,1271,1272,1273,1274,1277,1278,1279,1284,1285,1286,1289,1290,1292,1295,1296,1298,1299,1302,1305,1319,1320,1321,1322,1324,1325,1326,1327,1347,1348,1349,1350,1354,1355,1356,1358,1359,1371,1372,1374,1403,1405,1418,1419,1421,1422,1433,1434,1436,1437,1440,1441,1442,1443,1445,1446,1451,1452,1453,1455,1456,1457,1461,1462,1486,1491,1504,1506,1509,1513,1519,1521,1525,1528,1531,1536,1538,1553,1558,1563,1570,1574,1578,1584,1586,1590,1592,1596,1598,1602,1606,1610,1612,1616,1618,1622,1624,1628,1630,1632])
    damaged={335,790,1097,1098,1100,1101,1144,1145,1146,1147,1149,1150,1284,1285,1286,1290,1302,1433,1442,1445,1446,1457,1461}
    for index in sorted(inline-damaged):
        line=by_index[index]['text']
        # Existing ASCII apostrophes inside English contractions remain intact.
        for match in re.finditer(r"([^‘’“”\";]+?)[‘“\"]((?:’(?=[A-Za-z])|[^‘’“”\"])+)",line):
            form=re.sub(r'^(?:[ABCD]:\s*|(?:Habitual|In progress):\s*)','',match[1].strip())
            if '/i/' in form or re.search(r'/[ɩʋəɛɔaeou]/',form): form=form.split('/',2)[-1].strip()
            form=re.sub(r'^\([^)]*\)\s*','',form)
            form=form.strip().strip(':').strip()
            if ':' in form:form=form.rsplit(':',1)[1].strip()
            # Alternate/optional forms are recorded individually, never a blended sentence.
            if re.search(r'[()/]|\bOR\b',form):
                if len(form.split())==1:
                    for part in re.split(r'[/()]',form):
                        if part:pair(part,match[2],index,note='Printed alternative form.')
                else:withheld.append({'sourceRef':ref(index),'form':form,'meaning':match[2],'reason':'Optional or damaged bundle; see source reference.'})
                continue
            form=clean(form)
            if len(form.split())==1 and re.match(r'^(?:I|We|They|She|He|You)\s+\w',match[2],re.I) and match[2].lower()!='you plural':
                withheld.append({'sourceRef':ref(index),'form':form,'meaning':match[2],'reason':'Sentence subject is missing in the transcription.'});continue
            pair(form,match[2],index,example=len(form.split())>1 or bool(re.search(r'[?!]',form)))
    # Clearly aligned interlinear cells give words and printed source glosses.
    interlinear=[294,297,300,306,308,314,316,324,328,330,349,351,354,356,365,369,374,416,432,457,459,461,476,478,493,495,497,629,634,697,709,712,733,739,743,747,751,754,760,771,785,808,819,822,833,837,844,848,867,1383,1387,1392,1396]
    for index in interlinear:
        row=by_index[index]['text']; following=by_index.get(index+1,{}).get('text','')
        ks=[clean(c) for c in row.split('\t') if c.strip()]; en=[clean(re.sub(r'\([SVOACTP©Q$£v]+\)|[+]|---+',' ',c)) for c in following.split('\t') if c.strip()]
        if len(ks)==len(en) and len(ks)>1:
            for form,meaning in zip(ks,en):pair(form,meaning,index,note='Meaning from an aligned source gloss; not a new translation.')
            sentence=' '.join(ks); translation=' '.join(en)
            for offset in (2,3):
                natural=by_index.get(index+offset,{}).get('text','').strip()
                if natural.startswith(('‘','“')):
                    translation=natural.strip('‘’“” ');break
            pair(sentence,translation,index,example=True,note='Printed example; aligned source rendering retained.',gloss=following)
    # Whole printed examples remain recoverable when individual gloss cells
    # are uneven. These renderings follow the source, without word harvesting
    # or composing new Kasem. Damaged renderings remain reference-only below.
    whole_renderings={297:'They laughed',306:'They beat us well',308:'They laughed much',
        349:'Thatone is my wife',351:'The man is brave',354:'She isnot my wife',356:'That is not the truth',
        369:'My wife is there',432:'I yesterday came',457:'He showed me',459:'We followed them',461:'They showed us',
        476:'It spoilec them',478:'They entered it',493:'I showed them',629:'room the is small',634:'House the is old',
        709:'The child ran away',733:'They got up anc ran off',739:'The child stro. g:y begged them',
        743:'they have already gone home',747:'Tam unabl>to go to market today',751:'I have finished eating',
        754:'We begged her father for a long time, without succe.',760:'I will never agree',
        771:'My inother went home and stayed at market',808:'when we got back home',
        822:'When we got home, we cooked some food',833:'to cook some food and eat it',
        837:'I went home to cook some food to eat',844:'so. that the food spoiled',
        848:'They beat us, and so we ran home',1383:'He went and finished reading a book',
        1387:'He will not go and finish reading a book',1392:'Then he went and read a book',1396:'Then he drank his fill of water'}
    for index,meaning in whole_renderings.items():
        pair(by_index[index]['text'],meaning,index,example=True,
             note='Whole printed example; source rendering retained despite uneven gloss columns.',
             gloss=by_index[index+1]['text'])
    # All noun paradigms with explicitly supplied singular/plural meanings.
    noun_indices=[912,913,914,915,918,919,920,921,924,925,926,932,933,934,935,938,939,940,941,942,943,944,951,952,953,956,957,958,959,960,961,962,963,964,965,968,969,970,971,972,979,981,982,983,986,987,990,991,992,993,996,997,998,999,1000,1007,1008,1009,1010,1012,1013,1014,1015,1016,1020,1021,1023,1026,1027,1030,1031,1032]
    for index in noun_indices:
        if index==1026:
            for form in ['badwon','badwoni','badwonna']:
                pair(form,'friend',index,'n.','Printed badwon(i) / badwonna paradigm; optional i retained as separate forms.')
            continue
        line=re.sub(r'^C-stems:\s*','',by_index[index]['text'])
        cells=[c.strip() for c in re.split(r'\t+|\s{2,}',line) if c.strip()]
        if len(cells)<2:withheld.append({'sourceRef':ref(index),'form':line,'meaning':'','reason':'Noun row not reliably aligned.'});continue
        meaning=cells[-1]; forms=cells[:-1]
        for number,cell in enumerate(forms):
            cell=re.sub(r'\(Class [A-E]\)','',cell).strip().lstrip('(')
            for form in cell.split('/'):
                form=re.sub(r'\b(?:wom|bam|dem|yam|kam|sem|kom|tem)\b','',form).strip()
                pair(form,meaning,index,'n.',f'Printed noun paradigm; {"singular" if number==0 else "plural"} form. The source gloss applies to the paradigm.')
    # All twelve tables are also preserved verbatim in the guide.
    for block in blocks:
        rows=block['rows'];index=block['index']
        if index==450:
            for row in rows[1:]:
                for cell in row[2:]:
                    m=re.match(r'(.+?)\s+([a-zɛɔŋáéó]+)$',cell,re.I)
                    if m:pair(m[2],m[1],index,'pronoun',f'{row[0]} {row[1]} personal pronoun, as printed.')
        elif index==487:
            for row in rows[1:]:
                for form in row[2].split('/'):
                    pair(form,row[3],index,'pronoun',f'Strong personal pronoun; {row[0]} {row[1]}.')
        elif index in [471,490,517,536]:
            for row in rows[1:]:
                for col in range(1 if index!=536 else 2,len(row)):
                    if row[col] and row[col] not in ['sg','pl']:
                        pair(row[col],rows[0][col],index,'pronoun' if index in [471,490] else 'determiner',f'Noun class {row[0]}, {row[1] if index==536 else rows[0][col]}; source-specific agreement form.')
        elif index==1472:
            for row in rows[1:]:
                meaning=row[5] or ('eat' if row[1]=='di' else '')
                for col in range(1,5):
                    for form in re.split(r'[ ()]+',row[col]):
                        if form:pair(form,meaning,index,'v.',f'Verb class {row[0]}; {rows[0][col]}. Parenthesized forms precede objects. Blank eat gloss is supported by Chapter 8 block 1200.')
    # Source-visible pairs where questions/phrases use tabs rather than quotation marks.
    for index in range(270,277):
        cells=[clean(c) for c in by_index[index]['text'].split('\t') if c.strip()]
        if len(cells)==2:pair(cells[0],cells[1],index,example=True)
    for index in range(279,289):
        cells=[clean(c) for c in by_index[index]['text'].split('\t') if c.strip()]
        if len(cells)==2:pair(cells[0],cells[1],index,example=len(cells[0].split())>1)
    # Clear greeting/gloss alignments whose labels and optional particles occupy cells.
    for form,meaning,index in [('De n lei','short greeting',198),('Ko gara','it is fine',221),('N na zaŋe to?','when you got up?',220),('Yazura wora','health is there',225),('Ko ye tɛ mo?','it is how?',226),('Dé ma daane daane','we pass time together',230),('N peini lanyerane','you sleep well',239),('N peini zurim','you sleep cool',239),('N ware soo','you greet house',248),('N yi lanyerane','you reach well',250),('Se jwa ne','until tomorrow',252),('Jwa ne daga','tomorrow is enough',252),('M biseim','you excuse us',265),('A zaaba','I welcome you-pl',265),('Dé zaane-m','we welcome you',260),('Dé zaaba','we welcome you-pl',262)]:
        pair(form,meaning,index,example=True,note='Greeting or printed greeting gloss; usage varies by situation and locality.')
    # Individually attested small words and question forms; never sentence-generated.
    for form,meaning,index in [('a','I',165),('na','when',166),('ba','do not',167),('wo','did not',168)]:
        pair(form,meaning,index,'grammatical form','Printed tone comparison; preserved without inferred accents.')
    pair('mɔne','laugh',1054,'v.','Clear second lexical pair; preceding beat gloss is damaged in the source.')
    pair('Ta!','Tell!',1461,example=True,note='Printed command; punctuation separates it from the following command.')
    pair('Zo!','Enter!',1461,example=True,note='Printed command; punctuation separates it from the preceding command.')
    for form,meaning,index in [('n','you',205),('ch','consonant digraph',172),('ny','consonant digraph',173),('ŋwaane','because of',174),('wora','to be there',367),('wo','is',367),('da','there',368),('wolo','who / which, noun class A singular',572),('balo','who / which, noun class A plural',572),('dedoa','one',550),('ma','then; next event in a sequence',1483),('maa','additional fact in continuous action',1488),('maa','motion in progress',1493),('daa','again',1508),('nam','emphasis or contrast; but',1511),('ya','past; previously',1370),('deem','remote past or previous occasion',1533),('doom','events in the present year',1540),('na','plural command particle',1551),('naa','attention particle used in greetings and instructions',1555),('mo','focus / attention to new information',1560),('yeini','where',1583),('yei','where',1583),('be','where',1583),('bɛ','what',1589),('wɔ','who',1595),('bera','who, plural',1595),('wɔ','whose',1601),('dɔ','which',1605),('kɔ','which',1605),('dɛ dɔ','when; which day',1609),('maŋa kɔ','when; which time',1609),('bagera','how many, noun class A',1615),('yagera','how many, noun class B',1615),('segera','how many, noun class C',1615),('tɛ','how',1621),('bɛ ŋwaane','why',1627),('bɛ mo ye','why',1627)]:
        pair(form,meaning,index,'grammatical form')
    # Restore English continuations only where the following printed line supplies them.
    continuations={699:"We haven’t yet gone again",1244:'She sweeps the room every day',1491:'And he was there for a long time',1506:'Then he got up and went home',1509:"He won’t go home again",1519:'He (had) wanted to go home',1521:'He had already gone home'}
    uncertain={349,354,390,476,478,586,599,610,615,621,634,733,739,747,754,771,798,815,833,844,1102,1131,1206,1220,1251,1252,1253,1271,1272,1295,1299,1321,1322,1371}
    for row in examples.values():
        index=int(row['sourceRef'].rsplit(' ',1)[1])
        if index in continuations:
            row['english']=continuations[index];row['note']+=' English continuation restored from the immediately following source line.'
        if index in uncertain or re.search(r'na ne|coiae|meaii|didnotsee|Weare|Iam|Iand',row['english']):
            row['ambiguous']=True;row['note']+=' Transcription uncertainty or optional meaning: reference only, excluded from exact translation answers.'
    payload={'importId':'gillbt-basic-grammar-1983-2014','sourceDocumentName':source.name,'sourceSha256':hashlib.sha256(source.read_bytes()).hexdigest(),'attribution':'P. L. Hewer, A Basic Grammar of Kasem, GILLBT. First printed 1983; supplied 2014 printing. Copyright GILLBT retained.','dialect':'Kasem (GILLBT grammar source)','sourceBlockCount':len(blocks),'tableCount':len(doc.tables),'chapters':[{'number':n,'title':t} for _,n,t in STARTS],'entries':list(entries.values()),'examples':list(examples.values()),'blocks':blocks,'withheld':withheld,'sourceIssues':['The editable copy contains transcription errors, missing words and punctuation, damaged column alignment and duplicated future-continuous material. Unrecoverable forms stay in the chapter reference and are excluded from exact answers.','Phonetic symbols are descriptions of sounds, not additional keyboard letters. The consonant-list count/list contains apparent transcription mistakes; the explicit ch/ny notes and seven written vowels agree with the orthography guide.','Source-specific spellings, tone marking and class labels are preserved. A pronoun table prints dé for class B singular, whereas the noun-class table prints de. No blanket change to historical dictionary words or existing spelling rules is made.','An aligned source gloss is labelled as a source rendering. No speaker confirmation, recording, invented interlinear gloss or new sentence is created.']}
    payload['figures']=figures
    with zipfile.ZipFile(source) as archive:
        payload['sourcePagination']=[{'part':name,'text':value} for name in archive.namelist()
            if re.match(r'word/(?:header|footer)\d+\.xml$',name)
            for value in [' '.join(n.text or '' for n in ET.fromstring(archive.read(name)).iter(W+'t'))] if value]
    (out/'book.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(json.dumps({'blocks':len(blocks),'tables':len(doc.tables),'figures':len(figures),'entries':len(entries),'examples':len(examples),'withheld':len(withheld)}))

if __name__=='__main__':main(sys.argv[1])
