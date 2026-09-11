from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
scenes = [
('When the sky was near', [
 ('Story title', 'Why the sky and clouds are far away', ''),
 ('Scene pill', 'Long ago', ''),
 ('Narration', 'Long, long ago, the land and the sky were close together. Clouds hung low over the compounds, within reach of an outstretched hand.', ''),
 ('Narration', 'At dusk, children gathered around an elder to hear how everything had changed.', ''),
 ('Speech bubble  Elder', 'Shall I tell you a story?', ''),
 ('Speech bubble  Children', 'Yes! We are listening!', ''),
]),
('Gathering clouds for cooking', [
 ('Scene pill', 'Clouds for the cooking pot', ''),
 ('Narration', 'People fetched clouds and cooked with them.', 'nɔɔna maa yeini ba kwoni konkweino tem ye ba mae ba saŋe'),
 ('Narration', 'When it was time to cook, people lifted their calabashes and gathered pieces of cloud. A little went into each cooking pot.', ''),
 ('Narration', 'A child reached for one more handful, but an elder smiled.', ''),
 ('Speech bubble  Elder', 'Gather only what we need for cooking.', ''),
 ('Speech bubble  Child', 'There is enough for everyone!', ''),
]),
('The old witch and her pestle', [
 ('Scene pill', 'An old witch', ''),
 ('Narration', 'Then an old witch began pounding fufu in her courtyard. Her pestle was unusually long.', ''),
 ('Narration', 'Each time she lifted the pestle, it hit the low sky. She frowned at the clouds above her.', ''),
 ('Sound effect', 'Thump, knock! Thump, knock!', ''),
 ('Speech bubble  Old witch', 'My fufu must be smooth!', ''),
 ('Word to use', 'fufu', 'sakɔra'),
]),
('The first warning', [
 ('Scene pill', 'The first warning', ''),
 ('Narration', 'A voice came from above, calm but clear. God warned the old witch.', ''),
 ('Speech bubble  God', 'Old woman, listen to me.', ''),
 ('Speech bubble  God', 'Stop hitting the sky.', 'ye ta n zɔ wɛ yuu kom'),
 ('Speech bubble  God', 'Lower your pestle. Each time you lift it so high, you strike the sky.', ''),
 ('Narration', 'She paused, glanced upward, and tightened her grip.', ''),
]),
('She would not listen', [
 ('Scene pill', 'Another warning', ''),
 ('Narration', 'But the old witch would not listen. She began pounding again, and her pestle struck the sky once more.', ''),
 ('Narration', 'God warned her again. Her neighbors hurried into the courtyard and begged her to stop.', ''),
 ('Speech bubble  Neighbors', 'Please listen. We all cook with these clouds.', ''),
 ('Speech bubble  Old witch', 'Let the sky move. I am not finished!', ''),
 ('Narration', 'She waved the neighbors away and raised her pestle again.', ''),
]),
('The final warning', [
 ('Scene pill', 'The final warning', ''),
 ('Narration', 'For a third time, God warned her. Everyone in the courtyard fell quiet.', ''),
 ('Speech bubble  God', 'If you keep striking the sky, I will raise it beyond your reach.', ''),
 ('Narration', 'The old witch lifted her pestle higher than before.', ''),
 ('Speech bubble  Old witch', 'Just one more blow!', ''),
 ('Sound effect', 'Knock!', ''),
]),
('The sky rises', [
 ('Scene pill', 'Up and away', ''),
 ('Narration', 'God made the sky rise.', 'wɛ maa pa wɛ yuu kom wo baŋa baŋa'),
 ('Narration', 'Up past the roofs, up past the tallest trees, up beyond every reaching hand. The clouds went with it.', ''),
 ('Speech bubble  Child', 'The clouds! They are going away!', ''),
 ('Narration', 'The villagers reached upward, but the clouds slipped farther away.', ''),
 ('Narration', 'The old witch lowered her pestle at last, but the sky kept rising.', ''),
]),
('Why the clouds are far away', [
 ('Scene pill', 'From that day onward', ''),
 ('Narration', 'That evening, the pots waited beside the hearth. Nobody could fetch a single cloud.', ''),
 ('Speech bubble  Child', 'Can we ever reach the clouds again?', ''),
 ('Narration', 'The elder looked at the distant sky before answering.', ''),
 ('Speech bubble  Elder', 'That is why we must listen when we are warned.', ''),
 ('Closing narration', 'And from that day, the sky and the clouds remained far above the land.', ''),
]),
]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.62)
sec.left_margin = sec.right_margin = Inches(.65)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title', 23), ('Heading 1', 17), ('Heading 2', 12)]:
 st = doc.styles[name]
 st.font.name = 'Arial'
 st.font.size = Pt(size)
 st.font.color.rgb = RGBColor(0,0,0)
 st.paragraph_format.space_after = Pt(10)

def fill(cell, text, label=None):
 cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
 p = cell.paragraphs[0]
 if label:
  r = p.add_run(label)
  r.bold = True
  r.font.size = Pt(9)
  p = cell.add_paragraph()
 if text:
  p.add_run(text)
 else:
  p.add_run('')
  p.paragraph_format.space_after = Pt(18)

for i, (title, rows) in enumerate(scenes):
 if i:
  doc.add_page_break()
 if i == 0:
  doc.add_paragraph('Sky folktale translation worksheet', 'Title')
  doc.add_paragraph('Translate the English into Navrongo Kasem in the right column. Your four contributions are already included exactly as written. Fill the blank cells and edit any wording you want to change.')
  doc.add_paragraph('Each numbered page is one illustrated scene. Scene pills are short captions; speech bubbles are character dialogue. Translate sound effects naturally or suggest a replacement.')
 doc.add_paragraph(f'Scene {i+1}  {title}', 'Heading 1')
 table = doc.add_table(rows=1, cols=2)
 table.autofit = False
 table.alignment = WD_TABLE_ALIGNMENT.CENTER
 table.columns[0].width = Inches(3.2)
 table.columns[1].width = Inches(3.77)
 pr = table._tbl.tblPr
 borders = OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e = OxmlElement('w:'+edge)
  for key,val in [('val','single'),('sz','5'),('color','D9D9D9')]: e.set(qn('w:'+key),val)
  borders.append(e)
 pr.append(borders)
 margins=OxmlElement('w:tblCellMar')
 for side in ['top','left','bottom','right']:
  e=OxmlElement('w:'+side);e.set(qn('w:w'),'110');e.set(qn('w:type'),'dxa');margins.append(e)
 pr.append(margins)
 for c, txt in zip(table.rows[0].cells,['English','Navrongo Kasem']):
  fill(c,txt)
  c.paragraphs[0].runs[0].bold=True
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF3');c._tc.get_or_add_tcPr().append(sh)
 repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
 for label,eng,kasem in rows:
  row=table.add_row()
  row.height=Inches(.65 if i==0 else .9)
  row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
  ns=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(ns)
  fill(row.cells[0],eng,label)
  fill(row.cells[1],kasem)
  for run in row.cells[1].paragraphs[0].runs:
   run.font.size=Pt(12)
 doc.add_paragraph('Translation notes or alternative wording', 'Heading 2')
 doc.add_paragraph('\n')

doc.core_properties.title='Sky folktale translation worksheet'
doc.core_properties.subject='English and Navrongo Kasem scene translation'
doc.core_properties.author=''
out=ROOT/'sky-folktale-translation.docx'
doc.save(out)
# Check editable text and exact user-supplied wording survived serialization.
check=Document(out)
text='\n'.join(c.text for t in check.tables for r in t.rows for c in r.cells)
for wording in ['sakɔra','ye ta n zɔ wɛ yuu kom','wɛ maa pa wɛ yuu kom wo baŋa baŋa','nɔɔna maa yeini ba kwoni konkweino tem ye ba mae ba saŋe']:
 assert wording in text, wording
assert len(check.tables)==8
print(out)
print(f'{out.stat().st_size} bytes; eight editable scene tables; all four user contributions verified')
