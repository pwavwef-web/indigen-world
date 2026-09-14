"""Build the illustrated publication from the user's completed Word worksheet."""
import argparse
import hashlib
import json
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase.pdfdoc import PDFString
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph
import pypdfium2 as pdfium
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[3]
OUT = REPO / 'output/pdf/sky-folktale'
W, H = 720, 1080
INK, PAPER, RUST = '#302c28', '#fff8ec', '#965534'

def extract(source):
    doc = Document(source)
    scenes, title = [], None
    for i, table in enumerate(doc.tables):
        rows = []
        for row in table.rows[1:]:
            left, kasem = [c.text.strip() for c in row.cells]
            label, english = left.split('\n', 1)
            label = label.strip()
            if not kasem:
                raise ValueError(f'Missing translation in scene {i+1}: {english}')
            item = dict(kind=label, english=english, kasem=kasem)
            if label == 'Story title': title = kasem
            else: rows.append(item)
        scenes.append(dict(id=i+1, rows=rows))
    assert len(scenes) == 8 and title
    return dict(title=title, language='Kasem', dialect='Navrongo', scenes=scenes,
                sourceSha256=hashlib.sha256(Path(source).read_bytes()).hexdigest(),
                translationSource='Completed translation worksheet supplied by the user',
                illustrationMethod='Eight separately generated illustrations; no crops from a montage')

def para(text, size=16, color=INK, bold=False, width=632):
    p = Paragraph(escape(text), ParagraphStyle('text', fontName='KasemBold' if bold else 'Kasem',
        fontSize=size, leading=size*1.35, textColor=HexColor(color)))
    _, height = p.wrap(width, H)
    return p, height

def put(c, text, x, y, width, size=16, color=INK, bold=False):
    p, height = para(text, size, color, bold, width)
    p.drawOn(c, x, y-height)
    return y-height

def main(source):
    data=extract(source)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT/'pages').mkdir(exist_ok=True)
    qa=REPO/'.tmp/sky-folktale-pdf-qa'
    qa.mkdir(parents=True, exist_ok=True)
    font_path=Path('C:/Windows/Fonts/arial.ttf')
    normal_font=TTFont('Kasem',str(font_path))
    cmap=normal_font.face.charToGlyph
    all_text=data['title']+''.join(r['kasem'] for s in data['scenes'] for r in s['rows'])
    assert not {ch for ch in all_text if not ch.isspace() and ord(ch) not in cmap}
    pdfmetrics.registerFont(normal_font)
    pdfmetrics.registerFont(TTFont('KasemBold','C:/Windows/Fonts/arialbd.ttf'))
    pdf=OUT/'sky-folktale-kasem.pdf'
    c=canvas.Canvas(str(pdf), pagesize=(W,H), pageCompression=1)
    c.setTitle(data['title']); c.setAuthor('Indigen World'); c.setSubject('Illustrated Navrongo Kasem folktale')
    c._doc.Catalog.Lang=PDFString('xsm')
    # Re-encode complete images for a compact PDF, preserving their entire frame.
    images={}
    for i in range(1,9):
        path=qa/f'scene-{i:02}.jpg'
        im=Image.open(ROOT/f'individual-scenes/scene-{i:02}.png').convert('RGB')
        im.save(path,quality=91,optimize=True)
        images[i]=path
    c.setFillColor(HexColor(PAPER));c.rect(0,0,W,H,fill=1,stroke=0)
    put(c,'INDIGEN WORLD  /  KASEM  /  NAVRONGO',44,1025,632,12,RUST,True)
    put(c,data['title'],44,963,632,38,INK,True)
    c.drawImage(str(images[7]),0,245,width=W,height=480)
    put(c,'Why the sky and clouds are far away',44,199,632,21)
    put(c,'A folktale in Navrongo Kasem',44,155,632,15,RUST)
    put(c,'Kasem translation supplied by the storyteller. Illustrations created with AI. Published by Indigen World.',44,76,632,10)
    c.showPage()
    for s in data['scenes']:
        c.setFillColor(HexColor(PAPER));c.rect(0,0,W,H,fill=1,stroke=0)
        pill=next(r['kasem'] for r in s['rows'] if r['kind']=='Scene pill')
        pill_width=min(628,pdfmetrics.stringWidth(pill,'KasemBold',17)+36)
        c.setFillColor(HexColor('#ead9bc'));c.roundRect(30,1013,pill_width,40,20,fill=1,stroke=0)
        put(c,pill,48,1041,pill_width-30,17,INK,True)
        c.drawImage(str(images[s['id']]),0,510,width=W,height=480)
        rows=[r for r in s['rows'] if r['kind'] not in ('Scene pill','Word to use')]
        # Measure before drawing; never shrink below a comfortable reading size.
        def measure(size):
            return sum(para(r['kasem'],size,width=550 if r['kind'].startswith('Speech') else 632)[1]
                +(32 if r['kind'].startswith('Speech') else 12) for r in rows)
        size=17
        while measure(size)>432 and size>15: size-=.5
        assert measure(size)<=432, (s['id'],measure(size))
        y=484
        for r in rows:
            if r['kind'].startswith('Speech'):
                right='Old witch' in r['kind'] or 'Child' in r['kind']
                x=92 if right else 44
                p, ph=para(r['kasem'],size,width=550)
                bh=ph+18
                bg='#e1e8ef' if right else '#f2e3c6'
                c.setFillColor(HexColor(bg));c.roundRect(x,y-bh,584,bh,14,fill=1,stroke=0)
                tail=c.beginPath()
                tx=x+555 if right else x+22
                tail.moveTo(tx,y-bh+3);tail.lineTo(tx+10,y-bh-8);tail.lineTo(tx+20,y-bh+3);tail.close()
                c.drawPath(tail,fill=1,stroke=0)
                p.drawOn(c,x+17,y-9-ph)
                y-=bh+14
            else:
                y=put(c,r['kasem'],44,y,632,size,RUST if r['kind']=='Sound effect' else INK,
                      r['kind']=='Sound effect')-12
        assert y>=42,(s['id'],y)
        put(c,f'{s["id"]} / 8',638,28,60,10,RUST)
        c.showPage()
    c.save()
    # Validate every translated publishing row is present, allowing layout whitespace.
    rendered=pdfium.PdfDocument(str(pdf))
    assert len(rendered)==9
    norm=lambda x:' '.join(x.split())
    extracted=norm(' '.join(p.extract_text() for p in PdfReader(pdf).pages))
    for s in data['scenes']:
        for r in s['rows']:
            if r['kind']!='Word to use': assert norm(r['kasem']) in extracted, r
    for i,p in enumerate(rendered):
        pix=p.render(scale=1.5).to_pil().convert('RGB')
        pix.save(qa/f'page-{i+1:02}.png')
        pix.save(OUT/f'pages/page-{i+1:02}.jpg',quality=91,optimize=True)
    # Full cover page, rather than an illustration crop.
    Image.open(OUT/'pages/page-01.jpg').save(OUT/'cover.jpg',quality=90,optimize=True)
    (ROOT/'story-published.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
    body='\n\n'.join(r['kasem'] for s in data['scenes'] for r in s['rows'] if r['kind']!='Word to use')
    record=dict(id='kasem-sky-far-away', title=data['title'],collectionKind='literature',
        publicationStatus='published',mediaType='document',language='Kasem',dialect='Navrongo',
        category='Folktale',description='An illustrated folktale in Navrongo Kasem.',
        englishSummary='Long ago, people fetched clouds for cooking. An old witch ignored repeated warnings and struck the sky with her pestle, so God raised the sky and clouds beyond reach.',
        body=body,creatorAttribution=dict(displayName='Indigen World'),
        licenceDisplay='Published with permission through Indigen World.',
        culturalNotes='Story and Navrongo Kasem translation supplied by the storyteller; illustrations created with AI.',
        pageCount=9,sourceSha256=data['sourceSha256'],schemaVersion=1)
    (OUT/'publication.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(dict(pdf=str(pdf),pages=len(rendered),bytes=pdf.stat().st_size,
                         translationRows=sum(len(s['rows']) for s in data['scenes'])),indent=2))

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('source');main(ap.parse_args().source)
